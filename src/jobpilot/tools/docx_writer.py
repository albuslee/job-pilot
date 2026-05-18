"""Render a TailoredCV into a .docx by mutating the user's template in-place.

Three regions are rewritten:
1. CAREER OVERVIEW: single paragraph between this anchor and SKILLS & EXPERTISE.
2. SKILLS & EXPERTISE: every body paragraph between this anchor and WORK EXPERIENCE.
3. Current-role `List Paragraph` bullets: only paragraphs whose style is
   `List Paragraph` between the role's role-line (identified by the caller-
   supplied anchor) and the following `Tech:` line. Any Normal-styled notes
   in that range (e.g. a "Promoted to ..." line) and the Tech: line itself
   pass through untouched.

If `current_role_anchor` is None or empty, region 3 is skipped — only the
summary and skills are rewritten. Everything else in the template (hackathon
sections, previous experience, education) is preserved verbatim. Missing
anchors raise SectionAnchorMissingError.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

from jobpilot.models.schemas import TailoredCV

_SUMMARY_ANCHOR = "CAREER OVERVIEW"
_SKILLS_ANCHOR = "SKILLS & EXPERTISE"
_WORK_ANCHOR = "WORK EXPERIENCE"
_TECH_PREFIX = "Tech:"
_LIST_STYLE = "List Paragraph"


class SectionAnchorMissingError(RuntimeError):
    """The template is missing a required section header or role anchor."""


def render_cv(
    *,
    template_path: Path,
    tailored: TailoredCV,
    resolved_current_role_bullets: list[str],
    output_path: Path,
    current_role_anchor: str | None = None,
) -> Path:
    doc = Document(str(template_path))

    _replace_section_block(
        doc,
        start_anchor=_SUMMARY_ANCHOR,
        end_anchor=_SKILLS_ANCHOR,
        new_lines=[tailored.summary],
    )
    _replace_section_block(
        doc,
        start_anchor=_SKILLS_ANCHOR,
        end_anchor=_WORK_ANCHOR,
        new_lines=list(tailored.skills_lines),
    )
    if current_role_anchor:
        _replace_current_role_bullets(doc, current_role_anchor, resolved_current_role_bullets)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _replace_section_block(
    doc: DocxDocument,
    *,
    start_anchor: str,
    end_anchor: str,
    new_lines: list[str],
) -> None:
    paragraphs = list(doc.paragraphs)
    start_idx = _find_exact(paragraphs, start_anchor)
    end_idx = _find_exact(paragraphs, end_anchor)
    body = paragraphs[start_idx + 1 : end_idx]
    _replace_body_paragraphs(body, new_lines, fallback_template=paragraphs[start_idx])


def _replace_current_role_bullets(
    doc: DocxDocument, role_anchor: str, new_bullets: list[str]
) -> None:
    paragraphs = list(doc.paragraphs)
    role_idx = _find_contains(paragraphs, role_anchor)
    tech_idx = _find_prefix_after(paragraphs, _TECH_PREFIX, role_idx + 1)
    bullet_paragraphs = [
        p for p in paragraphs[role_idx + 1 : tech_idx] if p.style and p.style.name == _LIST_STYLE
    ]
    if not bullet_paragraphs:
        raise SectionAnchorMissingError(
            f"No List Paragraph rows found between role anchor {role_anchor!r} and Tech: line."
        )
    _replace_body_paragraphs(bullet_paragraphs, new_bullets, fallback_template=bullet_paragraphs[0])


def _find_exact(paragraphs: list[Paragraph], anchor: str) -> int:
    for i, p in enumerate(paragraphs):
        if p.text.strip() == anchor:
            return i
    raise SectionAnchorMissingError(f"Anchor not found: {anchor!r}")


def _find_contains(paragraphs: list[Paragraph], needle: str) -> int:
    for i, p in enumerate(paragraphs):
        if needle in p.text:
            return i
    raise SectionAnchorMissingError(f"Anchor (contains) not found: {needle!r}")


def _find_prefix_after(paragraphs: list[Paragraph], prefix: str, start: int) -> int:
    for i in range(start, len(paragraphs)):
        if paragraphs[i].text.strip().startswith(prefix):
            return i
    raise SectionAnchorMissingError(f"Anchor (startswith) not found from {start}: {prefix!r}")


def _replace_body_paragraphs(
    body: list[Paragraph],
    new_lines: list[str],
    *,
    fallback_template: Paragraph,
) -> None:
    """Reuse existing paragraphs where possible, insert clones for overflow, delete extras."""
    template_p = body[0] if body else fallback_template

    last_anchor: Paragraph = template_p
    for i, line in enumerate(new_lines):
        if i < len(body):
            _set_paragraph_text(body[i], line)
            last_anchor = body[i]
        else:
            new_p = _insert_paragraph_after(last_anchor, line, template=template_p)
            body.append(new_p)
            last_anchor = new_p

    for extra in body[len(new_lines) :]:
        _remove_paragraph(extra)


def _set_paragraph_text(p: Paragraph, text: str) -> None:
    """Replace paragraph text while keeping the first run's formatting."""
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for run in p.runs[1:]:
        run._element.getparent().remove(run._element)


def _remove_paragraph(p: Paragraph) -> None:
    p._element.getparent().remove(p._element)


def _insert_paragraph_after(reference_p: Paragraph, text: str, *, template: Paragraph) -> Paragraph:
    """Insert a paragraph cloned from `template` right after `reference_p`."""
    new_element = deepcopy(template._element)
    reference_p._element.addnext(new_element)
    new_p = Paragraph(new_element, template._parent)
    _set_paragraph_text(new_p, text)
    return new_p
