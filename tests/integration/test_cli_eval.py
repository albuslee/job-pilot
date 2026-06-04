from __future__ import annotations

from pathlib import Path
from typing import Any
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from typer.testing import CliRunner

from jobpilot.cli import app
from jobpilot.models.schemas import EvaluationResult


@pytest.fixture
def cli_env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    profile = tmp_path / "profile"
    profile.mkdir()
    doc = Document()
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Built RAG pipelines with Anthropic.")
    doc.save(profile / "cv.docx")

    jd = tmp_path / "canva.txt"
    jd.write_text("Senior fullstack at Canva. Python, TypeScript, RAG a plus.")

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    monkeypatch.setenv("EMBEDDING_PROVIDER", "voyage")
    monkeypatch.setenv("VOYAGE_API_KEY", "vy-test")
    monkeypatch.setenv("PROFILE_DIR", str(profile))
    monkeypatch.setenv("CHROMA_DIR", str(tmp_path / "chroma"))
    return jd


def _fake_embedder() -> Any:
    class _Stub:
        def embed(self, texts: list[str]) -> list[list[float]]:
            return [[float(len(t)), 0.0] for t in texts]

    return _Stub()


def test_cli_version() -> None:
    runner = CliRunner()
    result = runner.invoke(app, ["version"])
    assert result.exit_code == 0
    assert "0.1.0" in result.stdout


def test_cli_ingest_then_eval(cli_env: Path) -> None:
    runner = CliRunner()

    expected_result = EvaluationResult(
        score=82,
        decision="apply",
        reasoning="Strong RAG + fullstack signals.",
        cited_chunk_ids=[],
        risk_flags=[],
    )
    mock_llm = MagicMock()
    mock_llm.with_structured_output.return_value.invoke.return_value = expected_result

    with (
        patch("jobpilot.cli.build_embedder", return_value=_fake_embedder()),
        patch("jobpilot.cli.build_llm", return_value=mock_llm),
    ):
        ingest = runner.invoke(app, ["ingest"])
        assert ingest.exit_code == 0, ingest.stdout
        assert "ingested" in ingest.stdout.lower()

        eval_ = runner.invoke(app, ["eval", str(cli_env)])
        assert eval_.exit_code == 0, eval_.stdout
        assert "82" in eval_.stdout
        assert "apply" in eval_.stdout.lower()
