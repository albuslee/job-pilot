from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

from docx import Document

from jobpilot.models.schemas import DocxSectionAssignment, DocxSectionNormalization
from jobpilot.rag.ingest import chunk_blocks, ingest_profile_dir
from jobpilot.rag.store import RagStore
from jobpilot.tools.docx_reader import DocxBlock


class _FakeEmbedder:
    def embed(self, texts: list[str]) -> list[list[float]]:
        return [[float(len(t)), 0.0] for t in texts]


class _CapturingStore:
    def __init__(self) -> None:
        self.chunks = []

    def add(self, chunks: list) -> None:
        self.chunks.extend(chunks)


def _write_cv(path: Path) -> None:
    doc = Document()
    doc.add_heading("Experience", level=1)
    doc.add_heading("ACME Corp", level=2)
    doc.add_paragraph("Built a RAG prototype.")
    doc.add_paragraph("Led the cloud networking migration.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, TypeScript, AWS.")
    doc.save(path)


def _write_complex_unstyled_cv(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Professional Summary")
    doc.add_paragraph("Senior Financial Accountant")
    doc.add_paragraph("Built regulated financial reporting systems.")
    doc.add_paragraph("Core Capabilities")
    doc.add_paragraph("Professional Skills: Financial reporting | Audit coordination")
    doc.add_paragraph("Professional Experience")
    doc.add_paragraph("Financial Accountant (External Reporting)")
    doc.add_paragraph("Chubb Insurance Australia Limited | Sydney | Jan 2024 - Dec 2025")
    doc.add_paragraph("Prepared statutory financial statements.")
    doc.add_paragraph("Key Achievements:")
    doc.add_paragraph("Improved balance sheet accuracy.")
    doc.add_paragraph("Financial Accountant")
    doc.add_paragraph("Cubic Transportation Systems | Sydney | Dec 2019 - Dec 2023")
    doc.add_paragraph("Delivered month-end close.")
    doc.add_paragraph("Qualifications")
    doc.add_paragraph("CPA Qualified")
    doc.save(path)


def test_chunk_blocks_groups_under_heading() -> None:
    blocks = [
        DocxBlock("Experience", ["Experience"], True, 1),
        DocxBlock("ACME Corp", ["Experience", "ACME Corp"], True, 2),
        DocxBlock("Built RAG prototype.", ["Experience", "ACME Corp"], False, 0),
        DocxBlock("Led migration.", ["Experience", "ACME Corp"], False, 0),
        DocxBlock("Skills", ["Skills"], True, 1),
        DocxBlock("Python.", ["Skills"], False, 0),
    ]
    chunks = chunk_blocks(blocks, source="cv.docx")
    headings = [c.heading_path for c in chunks]
    assert ["Experience", "ACME Corp"] in headings
    assert ["Skills"] in headings
    # Two body paragraphs under "ACME Corp" must be merged into the same chunk's text.
    tm = next(c for c in chunks if c.heading_path == ["Experience", "ACME Corp"])
    assert "Built RAG prototype." in tm.text
    assert "Led migration." in tm.text


def test_ingest_profile_dir_persists_chunks(tmp_path: Path) -> None:
    profile = tmp_path / "profile"
    profile.mkdir()
    _write_cv(profile / "cv.docx")

    store = RagStore(persist_dir=tmp_path / "chroma", embedder=_FakeEmbedder())
    n = ingest_profile_dir(profile, store=store)

    assert n >= 2  # at least Experience>ACME Corp and Skills
    hits = store.query("anything", k=10)
    assert any("RAG prototype" in h.text for h in hits)


def test_ingest_profile_dir_ignores_office_lock_docx_files(tmp_path: Path) -> None:
    profile = tmp_path / "profile"
    profile.mkdir()
    _write_cv(profile / "cv.docx")
    (profile / "~$cv.docx").write_bytes(b"not a real docx")
    store = RagStore(
        persist_dir=tmp_path / "chroma",
        collection="profile",
        embedder=_FakeEmbedder(),
    )

    n = ingest_profile_dir(profile, store=store)

    assert n >= 2
    hits = store.query("anything", k=10)
    assert all(hit.source == "cv.docx" for hit in hits)


def test_ingest_smart_docx_uses_llm_fallback_for_low_confidence_docx(
    tmp_path: Path,
) -> None:
    profile = tmp_path / "profile"
    profile.mkdir()
    doc = Document()
    doc.add_paragraph("Built production cloud systems.")
    doc.add_paragraph("Python, TypeScript, AWS.")
    doc.save(profile / "cv.docx")

    llm = MagicMock()
    llm.with_structured_output.return_value.invoke.return_value = DocxSectionNormalization(
        assignments=[
            DocxSectionAssignment(block_index=0, heading_path=["Summary"]),
            DocxSectionAssignment(block_index=1, heading_path=["Skills"]),
        ]
    )
    store = _CapturingStore()

    n = ingest_profile_dir(profile, store=store, smart_docx=True, llm=llm)

    assert n == 2
    llm.with_structured_output.assert_called_once()
    assert [c.heading_path for c in store.chunks] == [["Summary"], ["Skills"]]


def test_ingest_smart_docx_uses_llm_fallback_for_complex_unstyled_docx(
    tmp_path: Path,
) -> None:
    profile = tmp_path / "profile"
    profile.mkdir()
    _write_complex_unstyled_cv(profile / "cv.docx")

    llm = MagicMock()
    llm.with_structured_output.return_value.invoke.return_value = DocxSectionNormalization(
        assignments=[
            DocxSectionAssignment(block_index=0, heading_path=[]),
            DocxSectionAssignment(block_index=1, heading_path=[]),
            DocxSectionAssignment(block_index=2, heading_path=["Summary"]),
            DocxSectionAssignment(block_index=3, heading_path=["Summary"]),
        ]
    )
    store = _CapturingStore()

    ingest_profile_dir(profile, store=store, smart_docx=True, llm=llm)

    llm.with_structured_output.assert_called_once()


def test_ingest_smart_docx_skips_llm_when_parse_confidence_is_high(
    tmp_path: Path,
) -> None:
    profile = tmp_path / "profile"
    profile.mkdir()
    _write_cv(profile / "cv.docx")
    llm = MagicMock()
    store = _CapturingStore()

    n = ingest_profile_dir(profile, store=store, smart_docx=True, llm=llm)

    assert n >= 2
    llm.complete_structured.assert_not_called()
