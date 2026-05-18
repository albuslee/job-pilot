from __future__ import annotations

from pathlib import Path

from docx import Document

from jobpilot.models.schemas import TailoredCV
from jobpilot.tools.docx_writer import (
    SectionAnchorMissingError,
    render_cv,
)
from tests.fixtures.cv_template import ROLE_ANCHOR


def _tailored() -> TailoredCV:
    return TailoredCV(
        target_company="Canva",
        target_role="Senior Fullstack Engineer",
        summary="Tailored summary line one. Tailored summary line two.",
        skills_lines=[
            "Languages: TypeScript, Python",
            "AWS Serverless: Lambda, Step Functions",
            "AI / Agents: LiteLLM, RAG pipelines",
        ],
        current_role_bullet_ids=["entry-x"],  # not used by docx_writer; resolver runs in CLI
    )


def test_render_cv_writes_output_file(tmp_path: Path, cv_template: Path) -> None:
    output = tmp_path / "out.docx"
    result = render_cv(
        template_path=cv_template,
        tailored=_tailored(),
        resolved_current_role_bullets=["Bullet alpha.", "Bullet beta."],
        output_path=output,
        current_role_anchor=ROLE_ANCHOR,
    )
    assert result == output
    assert output.is_file()
    assert output.stat().st_size > 1000


def test_render_cv_replaces_summary(tmp_path: Path, cv_template: Path) -> None:
    output = tmp_path / "out.docx"
    render_cv(
        template_path=cv_template,
        tailored=_tailored(),
        resolved_current_role_bullets=["x"],
        output_path=output,
        current_role_anchor=ROLE_ANCHOR,
    )

    doc = Document(str(output))
    texts = [p.text.strip() for p in doc.paragraphs]
    overview_idx = texts.index("CAREER OVERVIEW")
    assert "Tailored summary line one" in texts[overview_idx + 1]


def test_render_cv_replaces_skills_in_order(tmp_path: Path, cv_template: Path) -> None:
    output = tmp_path / "out.docx"
    render_cv(
        template_path=cv_template,
        tailored=_tailored(),
        resolved_current_role_bullets=["x"],
        output_path=output,
        current_role_anchor=ROLE_ANCHOR,
    )

    doc = Document(str(output))
    texts = [p.text.strip() for p in doc.paragraphs]
    skills_idx = texts.index("SKILLS & EXPERTISE")
    work_idx = texts.index("WORK EXPERIENCE")
    skills_lines = [t for t in texts[skills_idx + 1 : work_idx] if t]
    assert skills_lines == [
        "Languages: TypeScript, Python",
        "AWS Serverless: Lambda, Step Functions",
        "AI / Agents: LiteLLM, RAG pipelines",
    ]


def test_render_cv_replaces_current_role_bullets(tmp_path: Path, cv_template: Path) -> None:
    output = tmp_path / "out.docx"
    new_bullets = [
        "POOL: Bullet one.",
        "POOL: Bullet two.",
        "POOL: Bullet three.",
    ]
    render_cv(
        template_path=cv_template,
        tailored=_tailored(),
        resolved_current_role_bullets=new_bullets,
        output_path=output,
        current_role_anchor=ROLE_ANCHOR,
    )

    doc = Document(str(output))
    paragraphs = list(doc.paragraphs)
    role_idx = next(i for i, p in enumerate(paragraphs) if ROLE_ANCHOR in p.text)
    tech_idx = next(
        i
        for i in range(role_idx + 1, len(paragraphs))
        if paragraphs[i].text.strip().startswith("Tech:")
    )
    bullets = [
        p.text.strip()
        for p in paragraphs[role_idx + 1 : tech_idx]
        if p.style and p.style.name == "List Paragraph"
    ]
    assert bullets == new_bullets


def test_render_cv_preserves_non_pool_paragraphs(tmp_path: Path, cv_template: Path) -> None:
    """Promoted-to-Staff note, Tech: line, and PREVIOUS EXPERIENCE must pass through verbatim."""
    output = tmp_path / "out.docx"
    render_cv(
        template_path=cv_template,
        tailored=_tailored(),
        resolved_current_role_bullets=["new bullet"],
        output_path=output,
        current_role_anchor=ROLE_ANCHOR,
    )

    src_texts = {p.text.strip() for p in Document(str(cv_template)).paragraphs}
    out_texts = {p.text.strip() for p in Document(str(output)).paragraphs}
    invariants = {
        t
        for t in src_texts
        if (
            "Promoted to" in t
            or t.startswith("Tech:")
            or "BetaCo" in t
            or "GammaWorks" in t
            or "DeltaCorp" in t
            or "Pioneer University" in t
        )
    }
    assert invariants, "fixture sanity: template must have known invariant blocks"
    assert invariants.issubset(out_texts), (
        f"missing invariants after render: {invariants - out_texts}"
    )


def test_render_cv_skips_role_bullets_when_anchor_is_none(
    tmp_path: Path, cv_template: Path
) -> None:
    """When no role anchor is configured, only summary + skills are rewritten."""
    output = tmp_path / "out.docx"
    render_cv(
        template_path=cv_template,
        tailored=_tailored(),
        resolved_current_role_bullets=[],
        output_path=output,
        current_role_anchor=None,
    )

    src_role_bullets = [
        p.text.strip()
        for p in Document(str(cv_template)).paragraphs
        if p.style and p.style.name == "List Paragraph"
    ]
    out_role_bullets = [
        p.text.strip()
        for p in Document(str(output)).paragraphs
        if p.style and p.style.name == "List Paragraph"
    ]
    assert src_role_bullets == out_role_bullets


def test_render_cv_raises_on_missing_anchor(tmp_path: Path) -> None:
    import pytest
    from docx import Document as Doc

    bad = tmp_path / "bad.docx"
    Doc().save(str(bad))

    output = tmp_path / "out.docx"
    with pytest.raises(SectionAnchorMissingError):
        render_cv(
            template_path=bad,
            tailored=_tailored(),
            resolved_current_role_bullets=["x"],
            output_path=output,
            current_role_anchor=ROLE_ANCHOR,
        )
