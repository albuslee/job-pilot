from __future__ import annotations

from pathlib import Path

from docx import Document

from jobpilot.tools.docx_reader import DocxBlock, read_docx, read_docx_detailed


def _build_fixture(path: Path) -> None:
    doc = Document()
    doc.add_heading("Experience", level=1)
    doc.add_heading("ACME Corp", level=2)
    doc.add_paragraph("Built RAG prototype.")
    doc.add_paragraph("Led cloud networking migration.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("BSc Computer Science.")
    doc.save(path)


def test_read_docx_yields_blocks_with_heading_path(tmp_path: Path) -> None:
    fixture = tmp_path / "cv.docx"
    _build_fixture(fixture)

    blocks = read_docx(fixture)

    assert all(isinstance(b, DocxBlock) for b in blocks)
    paragraphs = [b for b in blocks if not b.is_heading]
    assert any(b.text == "Built RAG prototype." for b in paragraphs)
    bullet = next(b for b in paragraphs if b.text == "Built RAG prototype.")
    assert bullet.heading_path == ["Experience", "ACME Corp"]
    edu = next(b for b in paragraphs if b.text == "BSc Computer Science.")
    assert edu.heading_path == ["Education"]


def test_read_docx_extracts_table_cells_in_document_order(tmp_path: Path) -> None:
    fixture = tmp_path / "cv.docx"
    doc = Document()
    doc.add_heading("Skills", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "Python, TypeScript"
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Built RAG prototype.")
    doc.save(fixture)

    blocks = read_docx(fixture)

    texts = [b.text for b in blocks]
    assert texts.index("Python, TypeScript") < texts.index("Experience")
    table_body = next(b for b in blocks if b.text == "Python, TypeScript")
    assert table_body.heading_path == ["Skills"]


def test_read_docx_detailed_detects_common_unstyled_section_names(tmp_path: Path) -> None:
    fixture = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Professional Summary")
    doc.add_paragraph("Built production cloud systems.")
    doc.add_paragraph("Work Experience")
    doc.add_paragraph("ACME Corp - Staff Engineer")
    doc.save(fixture)

    result = read_docx_detailed(fixture)

    summary = next(b for b in result.blocks if b.text == "Built production cloud systems.")
    experience = next(b for b in result.blocks if b.text == "ACME Corp - Staff Engineer")
    assert summary.heading_path == ["Professional Summary"]
    assert experience.heading_path == ["Work Experience"]
    assert result.diagnostics.heading_count == 2
    assert not result.diagnostics.fallback_recommended


def test_read_docx_detailed_recommends_fallback_for_complex_unstyled_cv(
    tmp_path: Path,
) -> None:
    fixture = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Professional Summary")
    doc.add_paragraph("Senior Financial Accountant")
    doc.add_paragraph("Built regulated financial reporting systems.")
    doc.add_paragraph("Core Capabilities")
    doc.add_paragraph("Professional Skills: Financial reporting | Audit coordination")
    doc.add_paragraph("Professional Experience")
    doc.add_paragraph("Financial Accountant (External Reporting)")
    doc.add_paragraph("Chubb Insurance Australia Limited | Sydney | Jan 2024 - Dec 2025")
    doc.add_paragraph("Prepared statutory financial statements.")
    doc.add_paragraph("Key Achievements:")
    doc.add_paragraph("Improved balance sheet accuracy.")
    doc.add_paragraph("Financial Accountant")
    doc.add_paragraph("Cubic Transportation Systems | Sydney | Dec 2019 - Dec 2023")
    doc.add_paragraph("Delivered month-end close.")
    doc.add_paragraph("Qualifications")
    doc.add_paragraph("CPA Qualified")
    doc.save(fixture)

    result = read_docx_detailed(fixture)

    headline = next(b for b in result.blocks if b.text == "Senior Financial Accountant")
    capabilities = next(b for b in result.blocks if b.text.startswith("Professional Skills:"))
    qualification = next(b for b in result.blocks if b.text == "CPA Qualified")
    role = next(b for b in result.blocks if b.text == "Financial Accountant (External Reporting)")

    assert headline.heading_path == ["Professional Summary"]
    assert capabilities.heading_path == ["Core Capabilities"]
    assert qualification.heading_path == ["Qualifications"]
    assert not role.is_heading
    assert result.diagnostics.styled_heading_count == 0
    assert result.diagnostics.inferred_heading_count >= 4
    assert result.diagnostics.fallback_recommended


def test_read_docx_detailed_recommends_fallback_for_low_confidence_parse(
    tmp_path: Path,
) -> None:
    fixture = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Built production cloud systems.")
    doc.add_paragraph("Python, TypeScript, AWS.")
    doc.save(fixture)

    result = read_docx_detailed(fixture)

    assert result.diagnostics.heading_count == 0
    assert result.diagnostics.assigned_body_ratio == 0.0
    assert result.diagnostics.fallback_recommended
