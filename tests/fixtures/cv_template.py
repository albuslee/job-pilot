"""Build a synthetic CV .docx that satisfies docx_writer's section anchors.

The real template (data/cv_template.docx) is gitignored because it contains the
user's actual CV content. Tests must not depend on that file existing, so we
build a minimal structurally-equivalent .docx at fixture time.

The structure mirrors a typical CV's anchors but uses generic content. Anchors
and the current-role List-Paragraph block are the only things docx_writer
cares about; the rest is filler to keep the preservation tests honest.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

# Generic placeholder anchors used throughout the test suite. These deliberately
# do NOT match any real employer; tests pass `ROLE_ANCHOR` into docx_writer.
ROLE_ANCHOR = "ACME Corp, Sydney"
ROLE_HEADER = f"{ROLE_ANCHOR} — Senior Engineer\tFebruary 2023 - Present"

_CURRENT_ROLE_BULLETS = [
    "Owned the X subsystem end-to-end across N regions.",
    "Reduced latency on the core path by Y% via Z.",
    "Drove a multi-quarter migration from A to B with zero downtime.",
    "Built tooling that automated a manual workflow used by N teams.",
    "Mentored junior engineers and ran cross-team design reviews.",
    "Partnered with product/design to ship the W feature on schedule.",
    "Adopted AI-assisted coding workflows; piloted them for the squad.",
]

_PROMOTED_NOTE = "Promoted to Senior Engineer, January 2024"


def build_minimal_template(path: Path) -> Path:
    """Write a fixture .docx to `path` and return it."""
    doc = Document()

    doc.add_paragraph("CAREER OVERVIEW")
    doc.add_paragraph("Original summary content from the template.")

    doc.add_paragraph("SKILLS & EXPERTISE")
    doc.add_paragraph("Languages: TypeScript, Python, JavaScript")
    doc.add_paragraph("AWS: Lambda, Step Functions, EventBridge")
    doc.add_paragraph("Data: PostgreSQL, DynamoDB, OpenSearch")

    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph(ROLE_HEADER)
    doc.add_paragraph(_PROMOTED_NOTE)
    for bullet in _CURRENT_ROLE_BULLETS:
        doc.add_paragraph(bullet, style="List Paragraph")
    doc.add_paragraph("Tech: TypeScript, Node.js, Python, AWS")

    doc.add_paragraph("PREVIOUS EXPERIENCE")
    doc.add_paragraph("BetaCo, City — Software Engineer\tMarch 2021 - January 2023")
    doc.add_paragraph("Built React frontends and Node.js services.", style="List Paragraph")
    doc.add_paragraph("GammaWorks, City — Full-Stack Engineer\tJune 2018 - March 2021")
    doc.add_paragraph(
        "Took ownership of the Delta product as lead developer.", style="List Paragraph"
    )
    doc.add_paragraph("DeltaCorp, City — Software Engineer\tDecember 2017 - May 2018")
    doc.add_paragraph("Built an Epsilon lookup product end-to-end.", style="List Paragraph")

    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("Pioneer University, City — Master of Engineering (IT)")

    doc.save(str(path))
    return path
